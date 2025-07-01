# --- START OF FILE app/routes.py ---

import json
from datetime import datetime
from io import BytesIO
from functools import wraps

from flask import (render_template, request, redirect, url_for, flash, jsonify,
                   send_file, make_response, Blueprint)
from flask_login import (login_user, logout_user, login_required, current_user)
from sqlalchemy import func, extract, or_
import openpyxl
from openpyxl import Workbook
from dateutil.relativedelta import relativedelta

from . import db
from .models import User, Record, PendingChange
from .utils import calculate_age_group

# Create a Blueprint. All routes are attached to this.
bp = Blueprint('routes', __name__, template_folder='templates')


# --- Custom Decorators ---
def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not current_user.is_authenticated or not current_user.is_admin():
            flash('Admin access is required for this page.', 'danger')
            return redirect(url_for('routes.dashboard'))
        return f(*args, **kwargs)
    return decorated_function


# --- Authentication Routes ---
@bp.route('/')
def index():
    if not current_user.is_authenticated:
        return redirect(url_for('routes.login'))
    return redirect(url_for('routes.dashboard'))

@bp.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('routes.dashboard'))

    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']
        user = User.query.filter_by(username=username).first()
        if user and user.check_password(password):
            login_user(user, remember=True)
            flash('تم تسجيل الدخول بنجاح', 'success')
            return redirect(url_for('routes.dashboard'))
        else:
            flash('اسم المستخدم أو كلمة المرور غير صحيحة', 'error')
    return render_template('login.html')

@bp.route('/logout')
@login_required
def logout():
    logout_user()
    flash('تم تسجيل الخروج بنجاح', 'info')
    return redirect(url_for('routes.login'))


# --- Main Application Routes ---
@bp.route('/dashboard')
@login_required
def dashboard():
    total_beneficiaries = Record.query.count()
    current_month = datetime.now().month
    current_year = datetime.now().year
    new_this_month = Record.query.filter(
        extract('month', Record.created_at) == current_month,
        extract('year', Record.created_at) == current_year
    ).count()
    needs_review = Record.query.filter_by(status='بحاجة لمراجعة').count()
    inactive_records = Record.query.filter_by(status='غير نشط').count()
    recent_beneficiaries = Record.query.order_by(Record.created_at.desc()).limit(5).all()
    
    pending_changes_count = 0
    if current_user.is_admin():
        pending_changes_count = PendingChange.query.filter_by(status='pending').count()
    
    stats = {
        'total_beneficiaries': total_beneficiaries, 'new_this_month': new_this_month,
        'needs_review': needs_review, 'inactive_records': inactive_records,
        'pending_changes_count': pending_changes_count
    }
    return render_template('dashboard.html', stats=stats, recent_beneficiaries=recent_beneficiaries)


# --- Beneficiary CRUD Routes ---
@bp.route('/beneficiaries')
@login_required
def beneficiaries():
    page = request.args.get('page', 1, type=int)
    per_page = 10
    query = Record.query
    search = request.args.get('search', '')
    if search:
        search_term = f"%{search}%"
        query = query.filter(or_(
            Record.first_name.ilike(search_term), Record.father_name.ilike(search_term),
            Record.grandfather_name.ilike(search_term), Record.family_name.ilike(search_term),
            Record.id_passport_number.ilike(search_term), Record.phone_number.ilike(search_term)
        ))
    status_filter = request.args.get('status', '')
    if status_filter:
        query = query.filter_by(status=status_filter)
    start_date_str = request.args.get('start_date', '')
    if start_date_str:
        query = query.filter(Record.created_at >= datetime.strptime(start_date_str, '%Y-%m-%d').date())
    end_date_str = request.args.get('end_date', '')
    if end_date_str:
        query = query.filter(Record.created_at <= datetime.strptime(end_date_str, '%Y-%m-%d').date())
    
    records = query.order_by(Record.created_at.desc()).paginate(page=page, per_page=per_page, error_out=False)
    return render_template('beneficiaries.html', records=records, search=search, 
                           status_filter=status_filter, start_date=start_date_str, end_date=end_date_str)

@bp.route('/add_single_beneficiary', methods=['GET', 'POST'])
@login_required
def add_single_beneficiary():
    if request.method == 'POST':
        try:
            record = Record(
                first_name=request.form['first_name'],
                father_name=request.form['father_name'],
                grandfather_name=request.form['grandfather_name'],
                family_name=request.form['family_name'],
                id_passport_number=request.form['id_passport_number'],
                date_of_birth=datetime.strptime(request.form['date_of_birth'], '%Y-%m-%d').date(),
                gender=request.form['gender'],
                marital_status=request.form['marital_status'],
                phone_number=request.form.get('phone_number', ''),
                family_members_count=int(request.form.get('family_members_count', 0)),
                address=request.form.get('address', ''),
                status=request.form['status'],
                created_by_user_id=current_user.id
            )
            db.session.add(record)
            db.session.commit()
            flash('تم إضافة المستفيد بنجاح', 'success')
            return redirect(url_for('routes.beneficiaries'))
        except Exception as e:
            db.session.rollback()
            flash(f'حدث خطأ أثناء إضافة المستفيد: {str(e)}', 'error')
    return render_template('add_single_beneficiary.html')

@bp.route('/edit_beneficiary/<int:record_id>', methods=['POST'])
@login_required
def edit_beneficiary(record_id):
    record = Record.query.get_or_404(record_id)
    try:
        updated_data = {
            'first_name': request.form['first_name'], 'father_name': request.form['father_name'],
            'grandfather_name': request.form['grandfather_name'], 'family_name': request.form['family_name'],
            'id_passport_number': request.form['id_passport_number'],
            'date_of_birth': request.form['date_of_birth'],
            'gender': request.form['gender'], 'marital_status': request.form['marital_status'],
            'phone_number': request.form.get('phone_number', ''),
            'family_members_count': int(request.form.get('family_members_count', 0)),
            'address': request.form.get('address', ''), 'status': request.form['status']
        }
        if current_user.is_admin():
            for key, value in updated_data.items():
                if key == 'date_of_birth': value = datetime.strptime(value, '%Y-%m-%d').date()
                setattr(record, key, value)
            record.updated_at = datetime.utcnow()
            db.session.commit()
            flash('تم تحديث بيانات المستفيد بنجاح', 'success')
        else:
            pending_change = PendingChange(
                record_id=record_id, user_id=current_user.id, change_type='update',
                changed_data=json.dumps(updated_data, default=str), status='pending'
            )
            db.session.add(pending_change)
            db.session.commit()
            flash('تم إرسال طلب التحديث للمراجعة', 'info')
    except Exception as e:
        db.session.rollback()
        flash(f'حدث خطأ أثناء تحديث البيانات: {str(e)}', 'error')
    return redirect(url_for('routes.beneficiaries'))

@bp.route('/delete_beneficiary/<int:record_id>', methods=['POST'])
@login_required
def delete_beneficiary(record_id):
    record = Record.query.get_or_404(record_id)
    try:
        if current_user.is_admin():
            db.session.delete(record)
            db.session.commit()
            flash('تم حذف المستفيد بنجاح', 'success')
        else:
            pending_change = PendingChange(
                record_id=record_id, user_id=current_user.id,
                change_type='delete', status='pending'
            )
            db.session.add(pending_change)
            db.session.commit()
            flash('تم إرسال طلب الحذف للمراجعة', 'info')
    except Exception as e:
        db.session.rollback()
        flash(f'حدث خطأ أثناء الحذف: {str(e)}', 'error')
    return redirect(url_for('routes.beneficiaries'))


# --- Approval Workflow Routes ---
@bp.route('/review_changes')
@admin_required
def review_changes():
    pending_changes = PendingChange.query.filter_by(status='pending').order_by(PendingChange.created_at.desc()).all()
    return render_template('review_changes.html', pending_changes=pending_changes)

@bp.route('/approve_change/<int:change_id>')
@admin_required
def approve_change(change_id):
    change = PendingChange.query.get_or_404(change_id)
    try:
        if change.change_type == 'update':
            record = Record.query.get(change.record_id)
            if record:
                updated_data = json.loads(change.changed_data)
                for key, value in updated_data.items():
                    if key == 'date_of_birth': value = datetime.strptime(value, '%Y-%m-%d').date()
                    setattr(record, key, value)
                record.updated_at = datetime.utcnow()
        elif change.change_type == 'delete':
            record = Record.query.get(change.record_id)
            if record: db.session.delete(record)
        change.status = 'approved'
        change.reviewed_by = current_user.id
        change.reviewed_at = datetime.utcnow()
        db.session.commit()
        flash('تم الموافقة على التغيير بنجاح', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'حدث خطأ أثناء الموافقة: {str(e)}', 'error')
    return redirect(url_for('routes.review_changes'))

@bp.route('/reject_change/<int:change_id>')
@admin_required
def reject_change(change_id):
    change = PendingChange.query.get_or_404(change_id)
    try:
        change.status = 'rejected'
        change.reviewed_by = current_user.id
        change.reviewed_at = datetime.utcnow()
        db.session.commit()
        flash('تم رفض التغيير', 'info')
    except Exception as e:
        db.session.rollback()
        flash(f'حدث خطأ أثناء الرفض: {str(e)}', 'error')
    return redirect(url_for('routes.review_changes'))


# --- Reports and Data Handling ---
@bp.route('/reports')
@login_required
def reports():
    gender_data = db.session.query(Record.gender, func.count(Record.id)).group_by(Record.gender).all()
    marital_data = db.session.query(Record.marital_status, func.count(Record.id)).group_by(Record.marital_status).all()
    status_data = db.session.query(Record.status, func.count(Record.id)).group_by(Record.status).all()
    
    records = Record.query.all()
    age_groups = {'0-18': 0, '19-35': 0, '36-50': 0, '51-65': 0, '65+': 0}
    for record in records:
        age_group = calculate_age_group(record.age)
        age_groups[age_group] += 1
        
    monthly_data = []
    for i in range(12):
        month_date = (datetime.now().replace(day=1) - relativedelta(months=i))
        count = Record.query.filter(extract('month', Record.created_at) == month_date.month,
                                    extract('year', Record.created_at) == month_date.year).count()
        monthly_data.append({'label': month_date.strftime('%Y-%m'), 'value': count})

    reports_data = {
        'gender': [{'label': item[0], 'value': item[1]} for item in gender_data],
        'marital_status': [{'label': item[0], 'value': item[1]} for item in marital_data],
        'status': [{'label': item[0], 'value': item[1]} for item in status_data],
        'age_groups': [{'label': k, 'value': v} for k, v in age_groups.items()],
        'monthly_growth': list(reversed(monthly_data))
    }
    return render_template('reports.html', reports_data=reports_data)

@bp.route('/import_export')
@login_required
def import_export():
    return render_template('import_export.html')

@bp.route('/download_template')
@login_required
def download_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Beneficiaries Template"
    headers = ['الاسم الأول', 'اسم الأب', 'اسم الجد', 'اسم العائلة', 'رقم الهوية/جواز السفر',
               'تاريخ الميلاد', 'الجنس', 'الحالة الاجتماعية', 'رقم الهاتف', 'عدد أفراد الأسرة', 'العنوان']
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    sample_data = ['أحمد', 'محمد', 'علي', 'الأحمد', '123456789', '1990-01-01',
                   'ذكر', 'متزوج', '0501234567', '4', 'الرياض']
    for col, data in enumerate(sample_data, 1):
        ws.cell(row=2, column=col, value=data)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='beneficiaries_template.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@bp.route('/export_records')
@login_required
def export_records():
    query = Record.query
    if request.args.get('type') == 'filtered':
        search = request.args.get('search', '')
        if search:
            search_term = f"%{search}%"
            query = query.filter(or_(
                Record.first_name.ilike(search_term), Record.father_name.ilike(search_term),
                Record.grandfather_name.ilike(search_term), Record.family_name.ilike(search_term),
                Record.id_passport_number.ilike(search_term), Record.phone_number.ilike(search_term)
            ))
        status_filter = request.args.get('status', '')
        if status_filter:
            query = query.filter_by(status=status_filter)
        start_date_str = request.args.get('start_date', '')
        if start_date_str:
            query = query.filter(Record.created_at >= datetime.strptime(start_date_str, '%Y-%m-%d').date())
        end_date_str = request.args.get('end_date', '')
        if end_date_str:
            query = query.filter(Record.created_at <= datetime.strptime(end_date_str, '%Y-%m-%d').date())

    records = query.order_by(Record.created_at.desc()).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Beneficiaries Export"
    headers = ['ID', 'الاسم الأول', 'اسم الأب', 'اسم الجد', 'اسم العائلة', 'رقم الهوية/جواز السفر',
               'تاريخ الميلاد', 'الجنس', 'الحالة الاجتماعية', 'رقم الهاتف', 'عدد أفراد الأسرة',
               'العنوان', 'الحالة', 'تاريخ الإنشاء']
    ws.append(headers)
    for record in records:
        ws.append([
            record.id, record.first_name, record.father_name, record.grandfather_name,
            record.family_name, record.id_passport_number,
            record.date_of_birth.strftime('%Y-%m-%d'), record.gender, record.marital_status,
            record.phone_number or '', record.family_members_count, record.address or '',
            record.status, record.created_at.strftime('%Y-%m-%d %H:%M')
        ])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'beneficiaries_export_{timestamp}.xlsx'
    return send_file(output, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@bp.route('/generate_pdf/<int:record_id>')
@login_required
def generate_pdf(record_id):
    record = Record.query.get_or_404(record_id)
    response_text = f"""
    تقرير المستفيد
    
    الاسم الكامل: {record.full_name}
    رقم الهوية: {record.id_passport_number}
    تاريخ الميلاد: {record.date_of_birth.strftime('%Y-%m-%d')}
    الجنس: {record.gender}
    الحالة الاجتماعية: {record.marital_status}
    رقم الهاتف: {record.phone_number or 'غير محدد'}
    عدد أفراد الأسرة: {record.family_members_count}
    العنوان: {record.address or 'غير محدد'}
    الحالة: {record.status}
    تاريخ الإنشاء: {record.created_at.strftime('%Y-%m-%d %H:%M')}
    """
    response = make_response(response_text)
    response.headers['Content-Type'] = 'text/plain; charset=utf-8'
    response.headers['Content-Disposition'] = f'attachment; filename="beneficiary_{record_id}.txt"'
    return response


# --- API Routes ---
@bp.route('/api/search_suggestions')
@login_required
def search_suggestions():
    query = request.args.get('q', '').strip()
    if len(query) < 2: return jsonify([])
    search_term = f"%{query}%"
    records = Record.query.filter(or_(
        Record.first_name.ilike(search_term), Record.father_name.ilike(search_term),
        Record.grandfather_name.ilike(search_term), Record.family_name.ilike(search_term),
        Record.id_passport_number.ilike(search_term), Record.phone_number.ilike(search_term)
    )).limit(5).all()
    suggestions = [{'name': r.full_name, 'id_passport': r.id_passport_number, 'phone': r.phone_number or '-'} for r in records]
    return jsonify(suggestions)

@bp.route('/api/statistics')
@login_required
def api_statistics():
    return jsonify({
        'total': Record.query.count(),
        'completed': Record.query.filter_by(status='مكتمل').count(),
        'review': Record.query.filter_by(status='بحاجة لمراجعة').count(),
        'inactive': Record.query.filter_by(status='غير نشط').count()
    })


# --- Settings and User Management ---
@bp.route('/settings')
@admin_required
def settings():
    users = User.query.order_by(User.created_at.desc()).all()
    return render_template('settings.html', users=users)

@bp.route('/add_user', methods=['POST'])
@admin_required
def add_user():
    if User.query.filter_by(username=request.form['username']).first():
        flash('اسم المستخدم موجود بالفعل', 'error')
        return redirect(url_for('routes.settings'))
    try:
        user = User(username=request.form['username'], full_name=request.form['full_name'], role=request.form['role'])
        user.set_password(request.form['password'])
        db.session.add(user)
        db.session.commit()
        flash('تم إضافة المستخدم بنجاح', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'حدث خطأ أثناء إضافة المستخدم: {str(e)}', 'error')
    return redirect(url_for('routes.settings'))

@bp.route('/edit_user/<int:user_id>', methods=['POST'])
@admin_required
def edit_user(user_id):
    user = User.query.get_or_404(user_id)
    if user.id == current_user.id and request.form['role'] != 'admin':
        flash('لا يمكنك تغيير دورك الإداري', 'error')
        return redirect(url_for('routes.settings'))
    try:
        if User.query.filter(User.username == request.form['username'], User.id != user_id).first():
            flash('اسم المستخدم موجود بالفعل', 'error')
            return redirect(url_for('routes.settings'))
        user.username = request.form['username']
        user.full_name = request.form['full_name']
        user.role = request.form['role']
        if request.form.get('password'):
            user.set_password(request.form['password'])
        db.session.commit()
        flash('تم تحديث بيانات المستخدم بنجاح', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'حدث خطأ أثناء تحديث البيانات: {str(e)}', 'error')
    return redirect(url_for('routes.settings'))

@bp.route('/delete_user/<int:user_id>', methods=['POST'])
@admin_required
def delete_user(user_id):
    user = User.query.get_or_404(user_id)
    if user.id == current_user.id:
        flash('لا يمكنك حذف حسابك الشخصي', 'error')
        return redirect(url_for('routes.settings'))
    if Record.query.filter_by(created_by_user_id=user_id).first():
        flash('لا يمكن حذف المستخدم لأنه قام بإنشاء سجلات في النظام', 'error')
        return redirect(url_for('routes.settings'))
    try:
        db.session.delete(user)
        db.session.commit()
        flash('تم حذف المستخدم بنجاح', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'حدث خطأ أثناء حذف المستخدم: {str(e)}', 'error')
    return redirect(url_for('routes.settings'))